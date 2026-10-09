"""Authorised source versions and bounded local document parsing."""

import base64
import binascii
import hashlib
import io
import zipfile
from contextlib import nullcontext
from datetime import UTC, datetime
from pathlib import PurePosixPath
from typing import Any
from uuid import UUID, uuid4

from sqlalchemy import Connection, Engine, delete, insert, select, update
from sqlalchemy.dialects.postgresql import insert as pg_insert

from memory_platform.auth.principal import Principal
from memory_platform.enums import Capability
from memory_platform.errors import AppError, ErrorCode
from memory_platform.knowledge_tables import knowledge_jobs, source_chunks, source_versions, sources
from memory_platform.storage import Storage
from memory_platform.tables import scopes

MAX_BYTES = 20 * 1024 * 1024
MAX_TEXT_BYTES = 20 * 1024 * 1024


def decode_upload(filename: str, encoded: str) -> tuple[str, bytes]:
    if PurePosixPath(filename).name != filename or "\\" in filename or "\x00" in filename:
        raise AppError(ErrorCode.validation_error, "Invalid source filename.")
    format_name = filename.rsplit(".", 1)[-1].lower()
    if format_name not in {"pdf", "md", "txt", "docx"}:
        raise AppError(ErrorCode.validation_error, "Unsupported source format.")
    try:
        value = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError):
        raise AppError(ErrorCode.validation_error, "Invalid base64 content.") from None
    if not value or len(value) > MAX_BYTES:
        raise AppError(ErrorCode.payload_too_large, "Source must contain 1 byte to 20 MiB.")
    return format_name, value


def parse_document(format_name: str, content: bytes) -> list[tuple[int | None, str]]:
    """Parse text only, never OCR, macros, or document instructions."""
    try:
        pages: list[tuple[int | None, str]]
        if format_name in {"txt", "md"}:
            pages = [(None, content.decode("utf-8-sig"))]
        elif format_name == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(content))
            if reader.is_encrypted:
                raise ValueError
            if len(reader.pages) > 300:
                raise AppError(ErrorCode.payload_too_large, "PDF exceeds 300 pages.")
            pages = [
                (number + 1, page.extract_text() or "") for number, page in enumerate(reader.pages)
            ]
        elif format_name == "docx":
            from docx import Document

            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                if sum(i.file_size for i in archive.infolist()) > MAX_TEXT_BYTES or any(
                    i.filename.endswith("vbaProject.bin") for i in archive.infolist()
                ):
                    raise AppError(ErrorCode.payload_too_large, "Document expansion exceeds limit.")
            document = Document(io.BytesIO(content))
            pages = [
                (
                    None,
                    "\n".join(
                        [p.text for p in document.paragraphs]
                        + [
                            "\t".join(c.text for c in r.cells)
                            for t in document.tables
                            for r in t.rows
                        ]
                    ),
                )
            ]
        else:
            raise ValueError
        if sum(len(t.encode("utf-8")) for _, t in pages) > MAX_TEXT_BYTES:
            raise AppError(ErrorCode.payload_too_large, "Extracted text exceeds limit.")
        if not any(text.strip() for _, text in pages):
            raise AppError(ErrorCode.validation_error, "No extractable text; OCR is not supported.")
        return pages
    except AppError:
        raise
    except Exception:
        raise AppError(ErrorCode.validation_error, "Document cannot be parsed.") from None


def chunk_pages(pages: list[tuple[int | None, str]]) -> list[tuple[int | None, str]]:
    chunks: list[tuple[int | None, str]] = []
    for page, text in pages:
        # Exact text slices keep source provenance; 200-character overlap preserves boundaries.
        for offset in range(0, len(text), 1800):
            part = text[offset : offset + 2000].strip()
            if part:
                chunks.append((page, part))
    if len(chunks) > 12000:
        raise AppError(ErrorCode.payload_too_large, "Too many document chunks.")
    return chunks


def checked_source(
    conn: Connection,
    principal: Principal,
    source_id: UUID,
    capability: Capability,
    *,
    lock: bool = False,
    include_deleted: bool = False,
) -> Any:
    query = select(sources).where(
        sources.c.id == source_id,
        sources.c.workspace_id == principal.workspace_id,
        *([] if include_deleted else [sources.c.deleted_at.is_(None)]),
    )
    if lock:
        query = query.with_for_update()
    row = conn.execute(query).mappings().first()
    if row is None:
        raise AppError(ErrorCode.not_found, "Not found.")
    principal.require(row["scope_id"], capability)
    return row


def enqueue_source(conn: Connection, workspace_id: UUID, version_id: UUID, version: int) -> UUID:
    job_id = uuid4()
    conn.execute(
        pg_insert(knowledge_jobs)
        .values(
            id=job_id,
            workspace_id=workspace_id,
            dedup_key=f"source:{version_id}",
            kind="source",
            target_id=version_id,
            target_version=version,
        )
        .on_conflict_do_nothing()
    )
    return job_id


def upload_source(
    engine: Engine,
    principal: Principal,
    storage: Storage,
    *,
    scope_id: UUID,
    title: str,
    filename: str,
    encoded: str | None = None,
    content: bytes | None = None,
    stored_version_id: UUID | None = None,
    connection: Connection | None = None,
    source_id: UUID | None = None,
    expected_version: int | None = None,
    settings: object = None,
) -> dict[str, Any]:
    principal.require(scope_id, Capability.source_ingest)
    format_name, content = decode_upload(
        filename, encoded if encoded is not None else base64.b64encode(content or b"").decode()
    )
    # Validate format/page limits before durable ingestion. Worker parses again after lease.
    parse_document(format_name, content)
    version_id = stored_version_id or uuid4()
    if stored_version_id is None:
        storage.put(str(version_id), content)
    try:
        with engine.begin() if connection is None else nullcontext(connection) as conn:
            from memory_platform.services.quotas import enforce_quota

            enforce_quota(
                conn,
                principal.workspace_id,
                settings,
                "sources",
                additional=1 if source_id is None else 0,
            )
            enforce_quota(
                conn, principal.workspace_id, settings, "source_bytes", additional=len(content)
            )
            if source_id is None:
                if (
                    conn.execute(
                        select(scopes.c.id).where(
                            scopes.c.id == scope_id, scopes.c.workspace_id == principal.workspace_id
                        )
                    ).first()
                    is None
                ):
                    raise AppError(ErrorCode.not_found, "Not found.")
                source_id = uuid4()
                version = 1
                conn.execute(
                    insert(sources).values(
                        id=source_id,
                        workspace_id=principal.workspace_id,
                        scope_id=scope_id,
                        title=title,
                        created_by=principal.actor_id,
                    )
                )
            else:
                row = checked_source(
                    conn, principal, source_id, Capability.source_ingest, lock=True
                )
                if row["scope_id"] != scope_id or row["current_version"] != expected_version:
                    raise AppError(ErrorCode.version_conflict, "Source version changed.")
                version = row["current_version"] + 1
                conn.execute(
                    update(sources).where(sources.c.id == source_id).values(current_version=version)
                )
                previous_ids = select(source_versions.c.id).where(
                    source_versions.c.source_id == source_id
                )
                conn.execute(
                    update(knowledge_jobs)
                    .where(
                        knowledge_jobs.c.target_id.in_(previous_ids),
                        knowledge_jobs.c.status.in_(["queued", "leased"]),
                    )
                    .values(status="cancelled", lease_token=None, lease_until=None)
                )
            conn.execute(
                insert(source_versions).values(
                    id=version_id,
                    source_id=source_id,
                    workspace_id=principal.workspace_id,
                    version=version,
                    filename=filename,
                    format=format_name,
                    content_hash=hashlib.sha256(content).hexdigest(),
                    storage_key=str(version_id),
                    byte_size=len(content),
                )
            )
            job_id = enqueue_source(conn, principal.workspace_id, version_id, version)
            conn.execute(
                update(scopes).where(scopes.c.id == scope_id).values(revision=scopes.c.revision + 1)
            )
        return {
            "id": source_id,
            "version_id": version_id,
            "version": version,
            "status": "queued",
            "job_id": job_id,
        }
    except Exception:
        if stored_version_id is None:
            storage.delete(str(version_id))
        raise


def delete_source(
    engine: Engine, principal: Principal, storage: Storage, source_id: UUID
) -> dict[str, Any]:
    with engine.begin() as conn:
        row = checked_source(
            conn,
            principal,
            source_id,
            Capability.memory_delete,
            lock=True,
            include_deleted=True,
        )
        versions = (
            conn.execute(
                select(source_versions.c.id, source_versions.c.storage_key).where(
                    source_versions.c.source_id == source_id
                )
            )
            .mappings()
            .all()
        )
        ids = [v["id"] for v in versions]
        conn.execute(
            update(sources).where(sources.c.id == source_id).values(deleted_at=datetime.now(UTC))
        )
        conn.execute(delete(source_chunks).where(source_chunks.c.source_version_id.in_(ids)))
        conn.execute(
            update(source_versions)
            .where(source_versions.c.id.in_(ids))
            .values(status="cancelled", summary=None)
        )
        conn.execute(
            update(knowledge_jobs)
            .where(knowledge_jobs.c.target_id.in_(ids))
            .values(status="cancelled", lease_token=None, lease_until=None)
        )
        conn.execute(
            update(scopes)
            .where(scopes.c.id == row["scope_id"])
            .values(revision=scopes.c.revision + 1)
        )
        # Record cleanup in the same transaction as concealment. Filesystem failure
        # or a crash after commit must not lose the obligation to erase originals.
        cleanup_id = conn.execute(
            pg_insert(knowledge_jobs)
            .values(
                workspace_id=principal.workspace_id,
                dedup_key=f"source-cleanup:{source_id}",
                kind="source_cleanup",
                target_id=source_id,
                target_version=row["current_version"],
            )
            .on_conflict_do_update(
                constraint="uq_knowledge_jobs_workspace_dedup",
                set_={
                    "status": "queued",
                    "attempts": 0,
                    "error_code": None,
                    "available_at": datetime.now(UTC),
                    "finished_at": None,
                    "lease_token": None,
                    "lease_until": None,
                },
                where=knowledge_jobs.c.status.in_(["failed", "cancelled"]),
            )
            .returning(knowledge_jobs.c.id)
        ).scalar_one_or_none()
    # Try immediate erasure as well. The worker owns durable retries; a DELETE
    # retry with current delete permission can requeue terminal cleanup failures.
    try:
        for version in versions:
            storage.delete(version["storage_key"])
    except OSError:
        pass
    else:
        if cleanup_id is not None:
            with engine.begin() as conn:
                conn.execute(
                    update(knowledge_jobs)
                    .where(
                        knowledge_jobs.c.id == cleanup_id,
                        knowledge_jobs.c.status == "queued",
                    )
                    .values(status="completed", finished_at=datetime.now(UTC))
                )
    return {"id": source_id, "deleted": True}
